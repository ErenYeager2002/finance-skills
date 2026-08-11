# -*- coding: utf-8 -*-
import json
import subprocess
import sys
from pathlib import Path

import openpyxl

import build_task_reports as B
import workbook_finalize as W


def _report(path, title, value):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = title
    ws.append(["项目", "值"])
    ws.append(["日期", value])
    wb.save(path)


def test_build_task_reports_creates_three_static_range_files(tmp_path):
    out = tmp_path / "04_产出"
    out.mkdir()
    for token, date in (("20260808", "2026-08-08"), ("20260809", "2026-08-09")):
        for prefix in ("核销日清", "变更清单", "订单写入差异"):
            _report(out / f"{prefix}_{token}.xlsx", "明细", date)
        (out / f"判定结果_{token}.json").write_text(
            json.dumps({
                "payment_count": 2,
                "counts": {"total": 3, "auto": 2, "hold": 1, "exception": 0},
            }, ensure_ascii=False),
            encoding="utf-8",
        )

    outputs = B.build(tmp_path, "2026-08-08", "2026-08-09")

    assert len(outputs) == 3
    for path in outputs:
        wb = openpyxl.load_workbook(path, data_only=False)
        assert wb.sheetnames[0] == "任务范围"
        assert {row[0].value for row in wb["任务范围"].iter_rows(min_row=2)} == {
            "2026-08-08", "2026-08-09"
        }
        assert len(wb.sheetnames) == 3
        wb.close()
        audit = W.inspect_calculation(path)
        assert audit.formula_cells == 0
        assert audit.full_calc_on_load == "0"
        assert audit.force_full_calc == "0"
        checker = Path(r"D:\BESTEASY\03_tools\scripts\xlsx_lightweight_audit.py")
        checked = subprocess.run(
            [sys.executable, str(checker), str(path), "--strict"],
            capture_output=True, text=True, encoding="utf-8",
        )
        assert checked.returncode == 0, checked.stdout + checked.stderr
